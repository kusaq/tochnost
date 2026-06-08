from datetime import datetime, timezone
import io

from fastapi import HTTPException, status
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill

from api.v1.base.service import BaseService
from api.v1.rail.fsm import detach_rail_from_sensor_state
from api.v1.rail.schemas import (
    RailsListResponse,
    RailUpdate,
    RailRead,
    RailMetricRead,
    SensorSeriesRequest,
    RailSensorSeries,
    SensorPoint,
    DeleteRailsResponse,
    RejectRailResponse,
)
from api.v1.stream_monitor.pipeline_hooks import emit_pipeline_event
from infra.timescale_db.models import ScrewStatus, RailStatus


class RailService(BaseService):
    async def list_rails(
        self,
        *,
        name: str | None,
        status: str | None,
        fastening_type: str | None,
        object_name: str | None,
        limit: int,
        offset: int,
        sort_by: str | None = None,
        order: str | None = None,
    ) -> RailsListResponse:
        items, total = await self.uow.rail.list_filtered_paginated(
            name=name,
            status=status,
            fastening_type=fastening_type,
            object_name=object_name,
            limit=limit,
            offset=offset,
            sort_by=(sort_by or "rail_id"),
            order_desc=(order or "desc").lower() != "asc",
        )
        return RailsListResponse(items=[RailRead.model_validate(i) for i in items], total=total)

    async def update_rail(self, rail_id: int, payload: RailUpdate) -> RailRead | None:
        fields = payload.model_dump(exclude_unset=True)
        if "scanned_name" in fields:
            raw = fields["scanned_name"]
            if raw is not None:
                stripped = raw.strip()
                fields["scanned_name"] = stripped if stripped else None
        updated = await self.uow.rail.update_fields(rail_id, **fields)
        return RailRead.model_validate(updated) if updated else None

    async def delete_rail(self, rail_id: int) -> None:
        rail = await self.uow.rail.get_by_id(rail_id)
        if rail is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Rail not found")
        await detach_rail_from_sensor_state(rail_id, self.redis)
        await self.uow.rail.delete_by_id(rail_id)

    async def delete_rails(self, rail_ids: list[int]) -> DeleteRailsResponse:
        not_found: list[int] = []
        for rail_id in rail_ids:
            rail = await self.uow.rail.get_by_id(rail_id)
            if rail is None:
                not_found.append(rail_id)
            else:
                await detach_rail_from_sensor_state(rail_id, self.redis)

        found_ids = [rid for rid in rail_ids if rid not in not_found]
        deleted_ids = await self.uow.rail.delete_by_ids(found_ids) if found_ids else []
        deleted = len(deleted_ids)
        if deleted == 0:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Rails not found")
        return DeleteRailsResponse(deleted=deleted, not_found=not_found)

    async def reject_rail(self, rail_id: int) -> RejectRailResponse:
        rail = await self.uow.rail.get_by_id(rail_id)
        if rail is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Rail not found")
        if rail.status != RailStatus.IN_PROGRESS:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail="Only in-progress rails can be rejected",
            )
        await detach_rail_from_sensor_state(rail_id, self.redis)
        await emit_pipeline_event(
            "rshr_discarded",
            rshr_id=rail_id,
            event_ts=datetime.now(timezone.utc),
            summary=f"РШР #{rail_id} отбракована вручную",
            payload={"reason": "manual_reject", "discarded": True},
        )
        await self.uow.rail.delete_by_id(rail_id)
        return RejectRailResponse(deleted=1, rail_id=rail_id)

    async def get_aggregated_metrics(self, rail_id: int) -> list[RailMetricRead]:
        # Получаем все Sensor1 и Sensor2 по рельсе
        s1_list = await self.uow.sensor1.list_by_rail(rail_id=rail_id, limit=100000, offset=0)
        s2_list = await self.uow.sensor2.list_by_rail(rail_id=rail_id)

        # Собираем значения по ключам
        def collect_s1(key: str) -> list[float]:
            vals: list[float] = []
            for s in s1_list:
                v = getattr(s, key, None)
                if v is not None:
                    try:
                        vals.append(float(v))
                    except Exception:
                        pass
            return vals

        def collect_s1_with_ts(key: str) -> list[tuple[datetime, float]]:
            vals: list[tuple[datetime, float]] = []
            for s in s1_list:
                v = getattr(s, key, None)
                if v is not None:
                    try:
                        vals.append((s.timestamp, float(v)))
                    except Exception:
                        pass
            return vals

        def latest_s1(key: str) -> tuple[float | None, "datetime | None"]:
            for s in reversed(s1_list):
                v = getattr(s, key, None)
                if v is not None:
                    try:
                        return float(v), s.timestamp
                    except Exception:
                        continue
            return None, None

        def collect_s1_bool(key: str) -> list[bool]:
            vals: list[bool] = []
            for s in s1_list:
                v = getattr(s, key, None)
                if v is not None:
                    vals.append(bool(v))
            return vals

        def latest_s1_bool(key: str) -> tuple[bool | None, datetime | None]:
            for s in reversed(s1_list):
                v = getattr(s, key, None)
                if v is not None:
                    return bool(v), s.timestamp
            return None, None

        def collect_s1_bool_with_ts(key: str) -> list[tuple[datetime, bool]]:
            vals: list[tuple[datetime, bool]] = []
            for s in s1_list:
                v = getattr(s, key, None)
                if v is not None:
                    vals.append((s.timestamp, bool(v)))
            return vals

        def collect_s2(key: str) -> list[float]:
            vals: list[float] = []
            for s in s2_list:
                v = getattr(s, key, None)
                if v is not None:
                    try:
                        vals.append(float(v))
                    except Exception:
                        pass
            return vals

        def collect_s2_with_ts(key: str) -> list[tuple[datetime, float]]:
            vals: list[tuple[datetime, float]] = []
            for s in s2_list:
                v = getattr(s, key, None)
                if v is not None:
                    try:
                        vals.append((s.timestamp, float(v)))
                    except Exception:
                        pass
            return vals

        def latest_s2(key: str) -> tuple[float | None, "datetime | None"]:
            for s in reversed(s2_list):
                v = getattr(s, key, None)
                if v is not None:
                    try:
                        return float(v), s.timestamp
                    except Exception:
                        continue
            return None, None

        def collect_s2_int(key: str) -> list[int]:
            vals: list[int] = []
            for s in s2_list:
                v = getattr(s, key, None)
                if v is not None:
                    try:
                        vals.append(int(v))
                    except Exception:
                        pass
            return vals

        def collect_s2_int_with_ts(key: str) -> list[tuple[datetime, int]]:
            vals: list[tuple[datetime, int]] = []
            for s in s2_list:
                v = getattr(s, key, None)
                if v is not None:
                    try:
                        vals.append((s.timestamp, int(v)))
                    except Exception:
                        pass
            return vals

        def latest_s2_int(key: str) -> tuple[int | None, "datetime | None"]:
            for s in reversed(s2_list):
                v = getattr(s, key, None)
                if v is not None:
                    try:
                        return int(v), s.timestamp
                    except Exception:
                        continue
            return None, None

        def avg(vals: list[float]) -> float | None:
            return (sum(vals) / len(vals)) if vals else None

        def max_val(vals: list[float]) -> float | None:
            return max(vals) if vals else None

        # Получаем пороги для "required"
        thresholds = {t.value: t async for t in _async_iter(await self.uow.threshold.list_all())}

        def required_for(key: str) -> str | None:
            th = thresholds.get(key)
            if not th:
                return None
            unit = f" {th.unit_of_measurement}" if th.unit_of_measurement else ""
            return f"от {th.min_value}{unit} до {th.max_value}{unit}"

        items: list[RailMetricRead] = []

        # 1) Пустое агрегирование
        empty_keys = [
            ("encoder1", "encoder1"),
            ("encoder2", "encoder2"),
            ("encoder3", "encoder3"),
            ("encoder4", "encoder4"),
        ]
        for disp, key in empty_keys:
            vals = collect_s1(key)
            latest_value, latest_ts = latest_s1(key)
            items.append(
                RailMetricRead(
                    rail_id=rail_id,
                    name=disp,
                    value=latest_value,
                    required=None,
                    values=collect_s1_with_ts(key),
                    note=None,
                )
            )

        laser_keys = [
            ("laserOnRailLeft", "laser_on_rail_left"),
            ("laserOnRailRight", "laser_on_rail_right"),
            ("laserOnTieLeft", "laser_on_tie_left"),
            ("laserOnTieRight", "laser_on_tie_right"),
        ]
        for disp, key in laser_keys:
            latest_value, latest_ts = latest_s1_bool(key)
            # Для совместимости value остаётся числом (0/1) либо None
            num_value: float | None
            if latest_value is None:
                num_value = None
            else:
                num_value = 1.0 if latest_value else 0.0
            items.append(
                RailMetricRead(
                    rail_id=rail_id,
                    name=disp,
                    value=num_value,
                    required=None,
                    values=collect_s1_bool_with_ts(key),
                    note=None,
                )
            )

        s2_status_keys = [
            ("Состояние ПЧ1", "frequency_status_1"),
            ("Состояние ПЧ2", "frequency_status_2"),
            ("Состояние ПЧ3", "frequency_status_3"),
            ("Состояние ПЧ4", "frequency_status_4"),
        ]
        for disp, key in s2_status_keys:
            latest_value, latest_ts = latest_s2_int(key)
            num_value: float | None = float(latest_value) if latest_value is not None else None
            items.append(
                RailMetricRead(
                    rail_id=rail_id,
                    name=disp,
                    value=num_value,
                    required=None,
                    values=collect_s2_int_with_ts(key),
                    note=None,
                )
            )

        # 2) Максимум
        mm_along_vals = collect_s1("mm_along_rail")
        max_mm = max_val(mm_along_vals)
        if max_mm is not None:
            # ищем timestamp первого вхождения максимума с конца (последнее измерение с макс. значением)
            for s in reversed(s1_list):
                v = getattr(s, "mm_along_rail", None)
                if v is not None and float(v) == float(max_mm):
                    # timestamp можно использовать в будущем, но сейчас не включаем в схему
                    break
        items.append(
            RailMetricRead(
                rail_id=rail_id,
                name="mmAlongRail",
                value=max_mm,
                required=None,
                values=collect_s1_with_ts("mm_along_rail"),
                note=None,
            )
        )

        # 3) Среднее
        avg_map = [
            ("mmBoltHeightLeftInner", "mm_bolt_height_left_inner", "mm_bolt_height_left_inner"),
            ("mmBoltHeightLeftOuter", "mm_bolt_height_left_outer", "mm_bolt_height_left_outer"),
            ("mmBoltHeightRightInner", "mm_bolt_height_right_inner", "mm_bolt_height_right_inner"),
            ("mmBoltHeightRightOuter", "mm_bolt_height_right_outer", "mm_bolt_height_right_outer"),
            ("mmGauge", "mm_gauge", "mm_gauge"),
            ("mmSideWearLeft", "mm_side_wear_left", "mm_side_wear_left"),
            ("mmSideWearRight", "mm_side_wear_right", "mm_side_wear_right"),
            ("mmVerticalWearLeft", "mm_vertical_wear_left", "mm_vertical_wear_left"),
            ("mmVerticalWearRight", "mm_vertical_wear_right", "mm_vertical_wear_right"),
            ("radRailTiltLeft", "rad_rail_tilt_left", "rad_rail_tilt_left"),
            ("radRailTiltRight", "rad_rail_tilt_right", "rad_rail_tilt_right"),
        ]
        for disp, key, th_key in avg_map:
            vals = collect_s1(key)
            avg_val = avg(vals)
            latest_value, latest_ts = latest_s1(key)
            items.append(
                RailMetricRead(
                    rail_id=rail_id,
                    name=disp,
                    value=avg_val,
                    required=required_for(th_key),
                    values=collect_s1_with_ts(key),
                    note=None,
                )
            )

        # Средние по Sensor2: температура, влажность и сопротивление
        temp_vals = collect_s2("temperature")
        items.append(
            RailMetricRead(
                rail_id=rail_id,
                name="temperature",
                value=avg(temp_vals),
                required=required_for("temperature"),
                values=collect_s2_with_ts("temperature"),
                note=None,
            )
        )
        hum_vals = collect_s2("humidity")
        items.append(
            RailMetricRead(
                rail_id=rail_id,
                name="humidity",
                value=avg(hum_vals),
                required=required_for("humidity"),
                values=collect_s2_with_ts("humidity"),
                note=None,
            )
        )
        res_vals = collect_s2("resistance")
        items.append(
            RailMetricRead(
                rail_id=rail_id,
                name="resistance",
                value=avg(res_vals),
                required=required_for("resistance"),
                values=collect_s2_with_ts("resistance"),
                note=None,
            )
        )

        torque_keys = [
            ("Момент ПЧ1", "frequency_torque_1"),
            ("Момент ПЧ2", "frequency_torque_2"),
            ("Момент ПЧ3", "frequency_torque_3"),
            ("Момент ПЧ4", "frequency_torque_4"),
        ]
        for disp, key in torque_keys:
            vals = collect_s2(key)
            items.append(
                RailMetricRead(
                    rail_id=rail_id,
                    name=disp,
                    value=avg(vals),
                    required=required_for("frequency_torque"),
                    values=collect_s2_with_ts(key),
                    note=None,
                )
            )

        conv_freq_keys = [
            ("Частота ПЧ1", "converter_frequency_1"),
            ("Частота ПЧ2", "converter_frequency_2"),
            ("Частота ПЧ3", "converter_frequency_3"),
            ("Частота ПЧ4", "converter_frequency_4"),
        ]
        for disp, key in conv_freq_keys:
            vals = collect_s2(key)
            items.append(
                RailMetricRead(
                    rail_id=rail_id,
                    name=disp,
                    value=avg(vals),
                    required=None,
                    values=collect_s2_with_ts(key),
                    note=None,
                )
            )

        return items

    async def export_metrics_excel(self, rail_id: int, metric_names: list[str] | None = None) -> bytes:
        """
        Формирует Excel-файл с агрегированными метриками по рельсе.
        Лист "Метрики": общая информация по рельсе + сводная таблица метрик.
        Отдельный лист на каждую метрику со списком исходных значений (timestamp, value),
        с подсветкой по threshold (если задан).
        
        Args:
            rail_id: Идентификатор рельсы
            metric_names: Список названий метрик для экспорта. Если None, экспортируются все метрики.
        """
        # Информация о рельсе
        rail = await self.uow.rail.get_by_id(rail_id)
        if rail is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Rail not found")

        screws = await self.uow.screw.list_by_rail(rail_id)
        tightened_count = sum(1 for s in screws if s.status == ScrewStatus.COMPLETED)

        # Пороговые значения
        thresholds_list = await self.uow.threshold.list_all()
        thresholds = {t.value: t async for t in _async_iter(thresholds_list)}

        # Карта: отображаемое имя метрики -> ключ threshold
        name_to_threshold_key: dict[str, str] = {
            "mmBoltHeightLeftInner": "mm_bolt_height_left_inner",
            "mmBoltHeightLeftOuter": "mm_bolt_height_left_outer",
            "mmBoltHeightRightInner": "mm_bolt_height_right_inner",
            "mmBoltHeightRightOuter": "mm_bolt_height_right_outer",
            "mmGauge": "mm_gauge",
            "mmSideWearLeft": "mm_side_wear_left",
            "mmSideWearRight": "mm_side_wear_right",
            "mmVerticalWearLeft": "mm_vertical_wear_left",
            "mmVerticalWearRight": "mm_vertical_wear_right",
            "radRailTiltLeft": "rad_rail_tilt_left",
            "radRailTiltRight": "rad_rail_tilt_right",
            "temperature": "temperature",
            "humidity": "humidity",
            "resistance": "resistance",
            "Момент ПЧ1": "frequency_torque",
            "Момент ПЧ2": "frequency_torque",
            "Момент ПЧ3": "frequency_torque",
            "Момент ПЧ4": "frequency_torque",
        }

        metrics = await self.get_aggregated_metrics(rail_id)
        
        # Фильтруем метрики, если указан список
        if metric_names is not None:
            metric_names_set = set(metric_names)
            metrics = [m for m in metrics if m.name in metric_names_set]

        wb = Workbook()
        ws = wb.active
        ws.title = "Метрики"

        # Общая информация о рельсе
        info_rows = [
            ("Rail ID", rail.rail_id),
            ("Название", rail.name or ""),
            ("Считанный номер", rail.scanned_name or ""),
            ("Статус", rail.status.value if getattr(rail, "status", None) is not None else ""),
            ("Сторона", rail.side.value if getattr(rail, "side", None) is not None else ""),
            ("Объект", rail.object_name or ""),
            ("Тип крепления", rail.fastening_type or ""),
            ("Шпал", rail.sleepers if rail.sleepers is not None else ""),
            ("Дата начала", rail.start_time),
            ("Дата окончания", rail.end_time),
            ("Всего закрученных гаек", tightened_count),
        ]
        for label, value in info_rows:
            ws.append([label, value])

        # Пустая строка перед сводной таблицей
        ws.append([])

        # Заголовок сводной таблицы
        headers = ["Название метрики", "Значение", "Требуемое значение", "Количество измерений"]
        ws.append(headers)

        # Данные сводной таблицы
        for m in metrics:
            ws.append(
                [
                    m.name,
                    m.value,
                    m.required or "",
                    len(m.values),
                ]
            )

        # Простое авто-расширение колонок на листе "Метрики"
        max_cols = ws.max_column
        for col_idx in range(1, max_cols + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 0
            for cell in ws[col_letter]:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max_len + 2

        # Отдельный лист для каждого списка значений
        used_titles: set[str] = {"Метрики"}
        for m in metrics:
            base_title = m.name[:25] if len(m.name) > 25 else m.name
            title = base_title or "Metric"
            suffix = 1
            while title in used_titles:
                suffix += 1
                # чтобы не превысить лимит Excel в 31 символ
                trimmed = base_title[: (31 - len(str(suffix)) - 1)]
                title = f"{trimmed}_{suffix}"
            used_titles.add(title)

            ws_metric = wb.create_sheet(title=title)

            th = thresholds.get(name_to_threshold_key.get(m.name, ""))
            if th:
                ws_metric.append(["threshold_min", th.min_value])
                ws_metric.append(["threshold_max", th.max_value])
                ws_metric.append([])  # пустая строка

            ws_metric.append(["timestamp", "value"])

            # Цвета заливки
            fill_ok = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")   # нежно-зелёный
            fill_bad = PatternFill(start_color="FDECEA", end_color="FDECEA", fill_type="solid")  # нежно-красный

            first_value_row = ws_metric.max_row + 1
            for ts, val in m.values:
                ws_metric.append([ts, val])

            # Подсветка по threshold (если есть)
            if th:
                min_v = th.min_value
                max_v = th.max_value
                for row_idx in range(first_value_row, ws_metric.max_row + 1):
                    cell = ws_metric.cell(row=row_idx, column=2)  # value
                    try:
                        v = float(cell.value)
                    except (TypeError, ValueError):
                        continue
                    if v < min_v or v > max_v:
                        fill = fill_bad
                    else:
                        fill = fill_ok
                    ws_metric.cell(row=row_idx, column=1).fill = fill
                    ws_metric.cell(row=row_idx, column=2).fill = fill

            # Авто-ширина колонок на листе метрики
            for col_idx in range(1, ws_metric.max_column + 1):
                col_letter = get_column_letter(col_idx)
                max_len = 0
                for cell in ws_metric[col_letter]:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
                ws_metric.column_dimensions[col_letter].width = max_len + 2

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()

    async def get_sensor_series(self, payload: SensorSeriesRequest) -> list[RailSensorSeries]:
        """
        Возвращает по каждой рельсе список точек:
        { rail_id, points: [ { timestamp, values{field: value} } ] }
        """
        result: list[RailSensorSeries] = []

        for rail_id in payload.rail_ids:
            if payload.source == "sensor1":
                rows = await self.uow.sensor1.list_by_rail(rail_id=rail_id, limit=100000, offset=0)
            else:
                rows = await self.uow.sensor2.list_by_rail(rail_id=rail_id)

            points: list[SensorPoint] = []
            for row in rows:
                values: dict[str, object] = {}
                for field in payload.fields:
                    if hasattr(row, field):
                        values[field] = getattr(row, field)
                if values:
                    points.append(SensorPoint(timestamp=row.timestamp, values=values))

            result.append(RailSensorSeries(rail_id=rail_id, points=points))

        return result


async def _async_iter(seq):
    for x in seq:
        yield x


